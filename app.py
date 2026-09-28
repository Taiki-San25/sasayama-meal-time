import base64
import binascii
import os
import re
import secrets
import uuid
from io import BytesIO
from urllib.parse import quote
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Literal

from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

import floor
import importer
from db import (ADMIN_ROLES, ENTRY_ROLES, IMPORT_ROLES, LAYOUT_ROLES, MEALS, FloorLayout, TableAssignment, ROLES, AuthLog, ChatMessage, ChatRead, Reservation, ReservationHistory, SessionLocal, TimeSlot,
                User, init_db, now_jst)
from security import hash_password, verify_password

BASE = Path(__file__).parent
IS_PROD = bool(os.environ.get("DATABASE_URL"))
ON_RENDER = bool(os.environ.get("RENDER"))  # Render が自動で設定する
TIME_RE = re.compile(r"^([01]\d|2[0-3]):[0-5]\d$")


def bootstrap_admin() -> None:
    """ユーザーが1人もいなければ最初の管理者を作る"""
    with SessionLocal() as s:
        if s.scalar(select(func.count()).select_from(User)):
            return
        username = os.environ.get("ADMIN_USERNAME")
        password = os.environ.get("ADMIN_PASSWORD")
        if not (username and password):
            if IS_PROD:
                print("[warn] ユーザー未登録: ADMIN_USERNAME / ADMIN_PASSWORD を設定して再起動してください")
                return
            username, password = "admin", "admin"  # ローカル開発用
            print("[info] ローカル開発用の管理者 admin / admin を作成しました")
        s.add(User(username=username, display_name="管理者", password_hash=hash_password(password),
                   role="admin", is_admin=True))
        s.commit()


def bootstrap_developer() -> None:
    """DEVELOPER_USERNAME / DEVELOPER_PASSWORD が設定され、そのIDが未登録なら最上位ロールのアカウントを作る"""
    username = os.environ.get("DEVELOPER_USERNAME", "").strip()
    password = os.environ.get("DEVELOPER_PASSWORD", "")
    if not (username and password):
        return
    with SessionLocal() as s:
        if s.scalar(select(User).where(User.username == username)):
            return
        s.add(User(username=username, display_name=ROLES["developer"], password_hash=hash_password(password),
                   role="developer", is_admin=True))
        s.commit()
        print(f"[info] 最上位ロールのアカウント {username} を作成しました")


@asynccontextmanager
async def lifespan(_app):
    if ON_RENDER and not IS_PROD:
        # 消えるディスク上の SQLite や開発用 admin/admin で本番起動しないよう止める
        raise RuntimeError("DATABASE_URL が未設定です。Render で PostgreSQL を接続してください")
    init_db()
    migrate_floor_layouts()
    bootstrap_admin()
    bootstrap_developer()
    yield


app = FastAPI(title="丹波篠山 喫食時間管理表", lifespan=lifespan)
app.add_middleware(
    SessionMiddleware,
    secret_key=os.environ.get("SESSION_SECRET") or secrets.token_hex(32),
    session_cookie="amt_session",
    max_age=60 * 60 * 12,
    same_site="lax",
    https_only=IS_PROD or ON_RENDER,
)
app.mount("/static", StaticFiles(directory=BASE / "static"), name="static")


@app.middleware("http")
async def no_stale_assets(request: Request, call_next):
    """デプロイ後に古い画面(JS/CSS/HTML)が残らないよう、毎回サーバーに更新を確認させる(未変更なら304)"""
    response = await call_next(request)
    if not request.url.path.startswith("/api/"):
        response.headers.setdefault("Cache-Control", "no-cache")
    return response


# ---------- 共通依存 ----------
def get_db():
    with SessionLocal() as s:
        yield s


def session_user(request: Request, db: Session) -> User | None:
    uid = request.session.get("uid")
    if uid is None:
        return None
    user = db.get(User, uid)
    return user if user and user.active else None


def current_user(request: Request, db: Session = Depends(get_db)) -> User:
    user = session_user(request, db)
    if not user:
        raise HTTPException(401, "ログインしてください")
    return user


def admin_user(user: User = Depends(current_user)) -> User:
    if user.role not in ADMIN_ROLES:
        raise HTTPException(403, "管理者のみ操作できます")
    return user


# ---------- 認証 ----------
class LoginIn(BaseModel):
    username: str
    password: str


@app.post("/api/login")
def login(body: LoginIn, request: Request, db: Session = Depends(get_db)):
    username = body.username.strip()
    user = db.scalar(select(User).where(User.username == username))
    if not user or not user.active or not verify_password(body.password, user.password_hash):
        db.add(AuthLog(user_id=user.id if user else None, username=username[:64], action="login_failed"))
        db.commit()
        raise HTTPException(401, "IDまたはパスワードが違います")
    db.add(AuthLog(user_id=user.id, username=user.username, action="login"))
    db.commit()
    request.session.clear()
    request.session["uid"] = user.id
    return {"ok": True}


@app.get("/logout")
def logout(request: Request, db: Session = Depends(get_db)):
    user = session_user(request, db)
    if user:
        db.add(AuthLog(user_id=user.id, username=user.username, action="logout"))
        db.commit()
    request.session.clear()
    return RedirectResponse("/login", 303)


@app.get("/api/me")
def me(user: User = Depends(current_user)):
    return {"id": user.id, "name": user.display_name, "username": user.username,
            "role": user.role, "role_label": ROLES.get(user.role, ""), "is_admin": user.role in ADMIN_ROLES}


class PasswordIn(BaseModel):
    current: str
    new: str = Field(min_length=4, max_length=128)


@app.post("/api/me/password")
def change_my_password(body: PasswordIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not verify_password(body.current, user.password_hash):
        raise HTTPException(400, "現在のパスワードが違います")
    user.password_hash = hash_password(body.new)
    db.add(AuthLog(user_id=user.id, username=user.username, action="password_change"))
    db.commit()
    return {"ok": True}


# ---------- 時間枠 ----------
def slot_labels(db: Session, meal: str) -> list[str]:
    return list(db.scalars(select(TimeSlot.label).where(TimeSlot.meal == meal).order_by(TimeSlot.sort)))


def check_meal(meal: str) -> str:
    if meal not in MEALS:
        raise HTTPException(404, "不明な区分です")
    return meal


@app.get("/api/slots/{meal}")
def get_slots(meal: str, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return slot_labels(db, check_meal(meal))


class SlotsIn(BaseModel):
    slots: list[str] = Field(max_length=48)

    @field_validator("slots")
    @classmethod
    def valid(cls, v):
        v = [s.strip() for s in v if s.strip()]
        for s in v:
            if not TIME_RE.match(s):
                raise ValueError(f"時刻は HH:MM 形式で入力してください: {s}")
        return sorted(set(v))


@app.put("/api/slots/{meal}")
def put_slots(meal: str, body: SlotsIn, _: User = Depends(admin_user), db: Session = Depends(get_db)):
    check_meal(meal)
    for ts in db.scalars(select(TimeSlot).where(TimeSlot.meal == meal)):
        db.delete(ts)
    db.add_all(TimeSlot(meal=meal, label=l, sort=i) for i, l in enumerate(body.slots))
    db.commit()
    return body.slots


# ---------- 台帳 ----------
def norm_time(v: str | None) -> str | None:
    if v in (None, ""):
        return None
    if not TIME_RE.match(v):
        raise ValueError("時刻は HH:MM 形式で入力してください")
    return v


class ReservationIn(BaseModel):
    """編集用(日付・泊数は変更しない)"""
    model_config = ConfigDict(str_strip_whitespace=True)

    room: str = Field(min_length=1, max_length=255)
    guest_name: str = Field(default="", max_length=128)
    adults: int = Field(default=0, ge=0, le=99)
    children: int = Field(default=0, ge=0, le=99)
    infants: int = Field(default=0, ge=0, le=99)
    time_slot: str | None = None
    allergy: str = Field(default="", max_length=2000)
    note: str = Field(default="", max_length=2000)
    grouped: bool = False           # グループ登録する
    group_with: int | None = None   # 紐づける相手の予約ID(None なら単体でグループにする)

    _time = field_validator("time_slot")(norm_time)


class ReservationCreateIn(ReservationIn):
    """登録用: date は管理表で表示中の日付。泊数分の日付に1件ずつ登録する"""
    date: date
    nights: int = Field(default=1, ge=1, le=30)


class TimeSlotIn(BaseModel):
    time_slot: str | None = None

    _time = field_validator("time_slot")(norm_time)


TRACKED = ("date", "nights", "night_no", "time_slot", "room", "guest_name", "adults", "children", "infants",
           "allergy", "note", "group_id", "entered_at")
GROUP_FIELDS = {"grouped", "group_with"}


def iso(v: datetime | None) -> str | None:
    return v.isoformat() if v else None


def snapshot(r: Reservation) -> dict:
    return {f: (v.isoformat() if isinstance(v, (date, datetime)) else v)
            for f in TRACKED for v in [getattr(r, f)]}


def to_dict(r: Reservation, names: dict[int, str]) -> dict:
    return {
        "id": r.id, "meal": r.meal, **snapshot(r),
        "created_at": iso(r.created_at), "created_by": names.get(r.created_by, ""),
        "updated_at": iso(r.updated_at), "updated_by": names.get(r.updated_by, ""),
        "deleted": r.deleted_at is not None,
        "entered_by": names.get(r.entered_by, ""),
        "deleted_at": iso(r.deleted_at), "deleted_by": names.get(r.deleted_by, ""),
        "resv_no": resv_no(r.ext_key),
    }


def resv_no(ext_key: str | None) -> str:
    """CSV取込元の予約番号。枝番0は省き、それ以外は「24472-2」"""
    if not ext_key:
        return ""
    no, _, branch = ext_key.partition("-")
    return no if branch in ("", "0") else ext_key


def user_names(db: Session) -> dict[int, str]:
    return dict(db.execute(select(User.id, User.display_name)).all())


def get_reservation(db: Session, meal: str, rid: int, *, editable: bool = True) -> Reservation:
    r = db.get(Reservation, rid)
    if not r or r.meal != meal:
        raise HTTPException(404, "見つかりません")
    if editable and r.deleted_at is not None:
        raise HTTPException(400, "削除済みの予約は変更できません")
    return r


def record(db: Session, r: Reservation, user: User, action: str, changes: dict) -> None:
    """更新日時・更新者を記録し、履歴を1件追記する"""
    now = now_jst()
    r.updated_at, r.updated_by = now, user.id
    db.flush()
    db.add(ReservationHistory(reservation_id=r.id, action=action, changes=changes, changed_at=now,
                              changed_by=user.id))


def change(db: Session, r: Reservation, values: dict, user: User, action: str = "update") -> None:
    """値を更新し、差分があれば履歴に残す(commit は呼び出し側)"""
    before = snapshot(r)
    for k, v in values.items():
        setattr(r, k, v)
    after = snapshot(r)
    diff = {f: [before[f], after[f]] for f in TRACKED if before[f] != after[f]}
    if diff:
        record(db, r, user, action, diff)


# ---------- グループ ----------
def group_target(db: Session, meal: str, d: date, target_id: int | None, self_id: int | None) -> Reservation:
    t = db.get(Reservation, target_id) if target_id else None
    if not t or t.meal != meal or t.date != d or t.deleted_at is not None or t.id == self_id:
        raise HTTPException(400, "紐づける予約を選んでください")
    return t


def ensure_group(db: Session, t: Reservation, user: User) -> str:
    """相手がまだグループでなければ新しいグループにする"""
    if not t.group_id:
        change(db, t, {"group_id": uuid.uuid4().hex}, user)
    return t.group_id


def shrink_group(db: Session, group_id: str | None, user: User) -> None:
    """抜けた結果1件だけ残ったグループは解消する"""
    if not group_id:
        return
    db.flush()
    rest = list(db.scalars(select(Reservation).where(Reservation.group_id == group_id,
                                                     Reservation.deleted_at.is_(None))))
    if len(rest) == 1:
        change(db, rest[0], {"group_id": None}, user)


@app.get("/api/{meal}/reservations")
def list_reservations(meal: str, d: date, include_deleted: bool = False, _: User = Depends(current_user),
                      db: Session = Depends(get_db)):
    check_meal(meal)
    q = select(Reservation).where(Reservation.meal == meal, Reservation.date == d)
    if not include_deleted:
        q = q.where(Reservation.deleted_at.is_(None))
    names = user_names(db)
    rows = [to_dict(r, names) for r in db.scalars(q.order_by(Reservation.room, Reservation.id))]
    if meal == "dinner":
        tables = reservation_tables(db, d)
        for r in rows:
            r["tables"] = tables.get(r["id"], [])
    return rows


@app.post("/api/{meal}/reservations")
def create_reservation(meal: str, body: ReservationCreateIn, user: User = Depends(current_user),
                       db: Session = Depends(get_db)):
    check_meal(meal)
    now = now_jst()
    values = body.model_dump(exclude={"date", "nights"} | GROUP_FIELDS)
    stay_id = uuid.uuid4().hex if body.nights > 1 else None
    target = group_target(db, meal, body.date, body.group_with, None) if body.grouped and body.group_with else None
    created = []
    for i in range(body.nights):
        d = body.date + timedelta(days=i)
        # 連泊時は、相手の同じ日の予約(相手も連泊なら)と紐づける
        mate = target if i == 0 or not target else (
            db.scalar(select(Reservation).where(Reservation.stay_id == target.stay_id, Reservation.date == d,
                                                Reservation.deleted_at.is_(None))) if target.stay_id else None)
        # 相手なしのグループ登録は単体のグループ(日ごとに別のID)
        group_id = ensure_group(db, mate, user) if mate else (uuid.uuid4().hex if body.grouped and not body.group_with else None)
        r = Reservation(meal=meal, date=d, nights=body.nights, night_no=i + 1,
                        stay_id=stay_id, group_id=group_id, **values,
                        created_at=now, created_by=user.id, updated_at=now, updated_by=user.id)
        db.add(r)
        db.flush()
        record(db, r, user, "create", {f: [None, v] for f, v in snapshot(r).items()})
        created.append(r)
    db.commit()
    names = user_names(db)
    return [to_dict(r, names) for r in created]


@app.put("/api/{meal}/reservations/{rid}")
def update_reservation(meal: str, rid: int, body: ReservationIn, user: User = Depends(current_user),
                       db: Session = Depends(get_db)):
    r = get_reservation(db, check_meal(meal), rid)
    values = body.model_dump(exclude=GROUP_FIELDS)
    old_group = r.group_id
    if not body.grouped:
        values["group_id"] = None
    elif not body.group_with:
        # 単体のグループ: すでに自分だけのグループならそのまま、他の予約と一緒なら抜けて単体にする
        alone = r.group_id and not db.scalar(select(func.count()).select_from(Reservation).where(
            Reservation.group_id == r.group_id, Reservation.id != r.id, Reservation.deleted_at.is_(None)))
        if not alone:
            values["group_id"] = uuid.uuid4().hex
    else:
        t = group_target(db, meal, r.date, body.group_with, r.id)
        if not (r.group_id and t.group_id == r.group_id):
            values["group_id"] = ensure_group(db, t, user)
    old_time = r.time_slot
    change(db, r, values, user)
    if old_time != r.time_slot:
        release_tables(db, r, user)
    if old_group and old_group != r.group_id:
        shrink_group(db, old_group, user)
    db.commit()
    return to_dict(r, user_names(db))


@app.patch("/api/{meal}/reservations/{rid}/time")
def set_time(meal: str, rid: int, body: TimeSlotIn, user: User = Depends(current_user),
             db: Session = Depends(get_db)):
    r = get_reservation(db, check_meal(meal), rid)
    old_time = r.time_slot
    change(db, r, {"time_slot": body.time_slot}, user)
    if old_time != r.time_slot:
        release_tables(db, r, user)
    db.commit()
    return to_dict(r, user_names(db))


class EnteredIn(BaseModel):
    entered: bool


@app.patch("/api/{meal}/reservations/{rid}/entered")
def set_entered(meal: str, rid: int, body: EnteredIn, user: User = Depends(current_user),
                db: Session = Depends(get_db)):
    """ステータス(入場済)の切り替え。レストランと最上位ロールのみ"""
    if user.role not in ENTRY_ROLES:
        raise HTTPException(403, "入場済の操作はレストランのみ可能です")
    r = get_reservation(db, check_meal(meal), rid)
    if body.entered != (r.entered_at is not None):
        change(db, r, {"entered_at": now_jst() if body.entered else None,
                       "entered_by": user.id if body.entered else None}, user)
        db.commit()
    return to_dict(r, user_names(db))


@app.delete("/api/{meal}/reservations/{rid}")
def delete_reservation(meal: str, rid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """論理削除: データは残し、削除日時・削除者を記録する"""
    r = get_reservation(db, check_meal(meal), rid)
    r.deleted_at, r.deleted_by = now_jst(), user.id
    record(db, r, user, "delete", {})
    release_tables(db, r, user)
    db.commit()
    return to_dict(r, user_names(db))


@app.post("/api/{meal}/reservations/{rid}/restore")
def restore_reservation(meal: str, rid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    r = get_reservation(db, check_meal(meal), rid, editable=False)
    if r.deleted_at is None:
        raise HTTPException(400, "削除されていない予約です")
    r.deleted_at, r.deleted_by = None, None
    record(db, r, user, "restore", {})
    db.commit()
    return to_dict(r, user_names(db))


@app.get("/api/{meal}/reservations/{rid}/history")
def reservation_history(meal: str, rid: int, _: User = Depends(current_user), db: Session = Depends(get_db)):
    get_reservation(db, check_meal(meal), rid, editable=False)
    names = user_names(db)
    rows = db.scalars(select(ReservationHistory).where(ReservationHistory.reservation_id == rid)
                      .order_by(ReservationHistory.id.desc()))
    return [{"action": h.action, "changes": h.changes, "changed_at": iso(h.changed_at),
             "changed_by": names.get(h.changed_by, "")} for h in rows]


# ---------- テーブルアサイン(夕食) ----------
class FloorTableIn(BaseModel):
    id: str = Field(pattern=r"^[A-Za-z0-9_-]{1,32}$")
    name: str = Field(min_length=1, max_length=16)
    seats: int = Field(ge=1, le=60)
    x: int = Field(ge=0, le=floor.CANVAS_W)
    y: int = Field(ge=0, le=floor.CANVAS_H)
    w: int = Field(ge=20, le=floor.CANVAS_W)
    h: int = Field(ge=20, le=floor.CANVAS_H)
    parts: list[dict] | None = None  # 連結前の卓(解除用)

    _name = field_validator("name")(lambda v: v.strip())


class FloorLayoutIn(BaseModel):
    tables: list[FloorTableIn] = Field(max_length=200)


class AssignIn(BaseModel):
    date: date
    time_slot: str
    table_id: str
    reservation_id: int


class UnassignIn(BaseModel):
    date: date
    time_slot: str
    table_id: str
    reservation_id: int | None = None  # 省略時はその卓の全予約


class MoveIn(UnassignIn):
    to_table_id: str


def layout_row(db: Session, d: date | None) -> FloorLayout | None:
    q = select(FloorLayout).where(FloorLayout.date.is_(None) if d is None else FloorLayout.date == d)
    return db.scalar(q)


def base_tables(db: Session) -> list[dict]:
    row = layout_row(db, None)
    return row.tables if row else floor.DEFAULT_TABLES


def effective_tables(db: Session, d: date) -> tuple[list[dict], bool]:
    """その日の配置(なければ基本レイアウト)と、その日専用の配置かどうか"""
    row = layout_row(db, d)
    return (row.tables, True) if row else (base_tables(db), False)


def save_layout(db: Session, d: date | None, tables: list[dict], user: User) -> None:
    row = layout_row(db, d)
    if not row:
        row = FloorLayout(date=d)
        db.add(row)
    row.tables, row.updated_at, row.updated_by = tables, now_jst(), user.id
    row.version = floor.LAYOUT_VERSION


def migrate_floor_layouts() -> None:
    """見取り図を変えたとき、保存済みの配置(基本・日ごと)を新しい見取り図に移し替える"""
    with SessionLocal() as s:
        for row in s.scalars(select(FloorLayout).where(FloorLayout.version < floor.LAYOUT_VERSION)):
            row.tables = floor.migrate_tables(row.tables, row.version, row.date is None)
            row.version = floor.LAYOUT_VERSION
            print(f"[info] テーブル配置({row.date or '基本'})を新しい見取り図に移し替えました")
        s.commit()


def table_names(db: Session, d: date) -> dict[str, str]:
    return {t["id"]: t["name"] for t in effective_tables(db, d)[0]}


def reservation_tables(db: Session, d: date) -> dict[int, list[str]]:
    """予約ID → 割り当てた卓名(配置の並び順)"""
    tables = effective_tables(db, d)[0]
    order = {t["id"]: i for i, t in enumerate(tables)}
    names = {t["id"]: t["name"] for t in tables}
    out: dict[int, list[tuple[int, str]]] = {}
    for a in db.scalars(select(TableAssignment).where(TableAssignment.date == d)):
        if a.table_id in names:
            out.setdefault(a.reservation_id, []).append((order[a.table_id], names[a.table_id]))
    return {rid: [n for _, n in sorted(v)] for rid, v in out.items()}


def record_tables(db: Session, r: Reservation, user: User, before: list[str]) -> None:
    after = reservation_tables(db, r.date).get(r.id, [])
    if before != after:
        record(db, r, user, "update", {"tables": [before, after]})


def release_tables(db: Session, r: Reservation, user: User) -> None:
    """予約の時間変更・削除時に、その予約のテーブル割り当てを外す"""
    if r.meal != "dinner":
        return
    rows = list(db.scalars(select(TableAssignment).where(TableAssignment.reservation_id == r.id)))
    if not rows:
        return
    before = reservation_tables(db, r.date).get(r.id, [])
    for a in rows:
        db.delete(a)
    db.flush()
    record_tables(db, r, user, before)


def layout_user(user: User = Depends(current_user)) -> User:
    if user.role not in LAYOUT_ROLES:
        raise HTTPException(403, "テーブル配置の編集権限がありません")
    return user


@app.get("/api/floor")
def floor_view(d: date, user: User = Depends(current_user), db: Session = Depends(get_db)):
    tables, own = effective_tables(db, d)
    names = user_names(db)
    rows = db.scalars(select(Reservation).where(Reservation.meal == "dinner", Reservation.date == d,
                                                Reservation.deleted_at.is_(None)))
    return {
        "date": d.isoformat(), "canvas": {"w": floor.CANVAS_W, "h": floor.CANVAS_H},
        "tables": tables, "own_layout": own, "can_edit_layout": user.role in LAYOUT_ROLES,
        "slots": slot_labels(db, "dinner"),
        "reservations": [to_dict(r, names) for r in rows],
        "assignments": [{"time_slot": a.time_slot, "table_id": a.table_id, "reservation_id": a.reservation_id}
                        for a in db.scalars(select(TableAssignment).where(TableAssignment.date == d))],
    }


@app.get("/api/floor/base")
def floor_base(_: User = Depends(current_user), db: Session = Depends(get_db)):
    return {"tables": base_tables(db), "canvas": {"w": floor.CANVAS_W, "h": floor.CANVAS_H}}


def check_layout(body: FloorLayoutIn) -> list[dict]:
    tables = [t.model_dump(exclude_none=True) for t in body.tables]
    ids = [t["id"] for t in tables]
    if len(set(ids)) != len(ids):
        raise HTTPException(400, "卓のIDが重複しています")
    seen = set()
    for t in tables:
        if t["name"] in seen:
            raise HTTPException(400, f"卓番号「{t['name']}」が重複しています")
        seen.add(t["name"])
    return tables


@app.put("/api/floor/base")
def put_floor_base(body: FloorLayoutIn, user: User = Depends(layout_user), db: Session = Depends(get_db)):
    save_layout(db, None, check_layout(body), user)
    db.commit()
    return {"tables": base_tables(db)}


@app.put("/api/floor/layout")
def put_floor_layout(d: date, body: FloorLayoutIn, user: User = Depends(layout_user),
                     db: Session = Depends(get_db)):
    """その日の配置を保存する。なくなった卓の割り当ては外す"""
    tables = check_layout(body)
    keep = {t["id"] for t in tables}
    removed = list(db.scalars(select(TableAssignment).where(TableAssignment.date == d,
                                                            TableAssignment.table_id.not_in(keep))))
    before = reservation_tables(db, d)
    for a in removed:
        db.delete(a)
    save_layout(db, d, tables, user)
    db.flush()
    after = reservation_tables(db, d)
    for rid in {a.reservation_id for a in removed}:
        r = db.get(Reservation, rid)
        if r and before.get(rid, []) != after.get(rid, []):
            record(db, r, user, "update", {"tables": [before.get(rid, []), after.get(rid, [])]})
    # 卓名の変更も割り当て済みの予約の履歴に残す
    for rid, names_ in after.items():
        if rid not in {a.reservation_id for a in removed} and before.get(rid, []) != names_:
            r = db.get(Reservation, rid)
            if r:
                record(db, r, user, "update", {"tables": [before.get(rid, []), names_]})
    db.commit()
    return {"tables": tables, "released": len(removed)}


@app.delete("/api/floor/layout")
def reset_floor_layout(d: date, user: User = Depends(layout_user), db: Session = Depends(get_db)):
    """その日の配置を基本レイアウトに戻す(基本にない卓の割り当ては外す)"""
    body = FloorLayoutIn(tables=[FloorTableIn(**t) for t in base_tables(db)])
    result = put_floor_layout(d, body, user, db)
    row = layout_row(db, d)
    if row:
        db.delete(row)
        db.commit()
    return result


def assign_target(db: Session, d: date, table_id: str) -> None:
    if table_id not in table_names(db, d):
        raise HTTPException(400, "テーブルが見つかりません")


def assigned(db: Session, d: date, time_slot: str, table_id: str, rid: int | None = None) -> list[TableAssignment]:
    q = select(TableAssignment).where(TableAssignment.date == d, TableAssignment.time_slot == time_slot,
                                      TableAssignment.table_id == table_id)
    if rid is not None:
        q = q.where(TableAssignment.reservation_id == rid)
    return list(db.scalars(q))


def freeze_layout(db: Session, d: date, user: User) -> None:
    """割り当てた日は、後で基本レイアウトを変えても崩れないようにその日の配置として保存する"""
    tables, own = effective_tables(db, d)
    if not own:
        save_layout(db, d, tables, user)


@app.post("/api/floor/assign")
def floor_assign(body: AssignIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    r = get_reservation(db, "dinner", body.reservation_id)
    if r.date != body.date or r.time_slot != body.time_slot:
        raise HTTPException(400, "予約の日付・時間が一致しません。画面を更新してください")
    assign_target(db, body.date, body.table_id)
    if assigned(db, body.date, body.time_slot, body.table_id, r.id):
        return {"ok": True}  # すでにこの卓に割り当て済み
    freeze_layout(db, body.date, user)
    before = reservation_tables(db, body.date).get(r.id, [])
    db.add(TableAssignment(date=body.date, time_slot=body.time_slot, table_id=body.table_id,
                           reservation_id=r.id, created_at=now_jst(), created_by=user.id))
    db.flush()
    record_tables(db, r, user, before)
    db.commit()
    return {"ok": True}


def find_assignments(db: Session, body: UnassignIn) -> list[TableAssignment]:
    rows = assigned(db, body.date, body.time_slot, body.table_id, body.reservation_id)
    if not rows:
        raise HTTPException(400, "割り当てが見つかりません。画面を更新してください")
    return rows


@app.post("/api/floor/unassign")
def floor_unassign(body: UnassignIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    rows = find_assignments(db, body)
    before = reservation_tables(db, body.date)
    for a in rows:
        db.delete(a)
    db.flush()
    for a in rows:
        record_tables(db, db.get(Reservation, a.reservation_id), user, before.get(a.reservation_id, []))
    db.commit()
    return {"ok": True}


@app.post("/api/floor/move")
def floor_move(body: MoveIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """卓の予約(reservation_id 省略時は全部)を別の卓へ移す。移動先に同じ予約があれば1つにまとめる"""
    rows = find_assignments(db, body)
    assign_target(db, body.date, body.to_table_id)
    if body.to_table_id == body.table_id:
        return {"ok": True}
    before = reservation_tables(db, body.date)
    for a in rows:
        if assigned(db, body.date, body.time_slot, body.to_table_id, a.reservation_id):
            db.delete(a)
        else:
            a.table_id = body.to_table_id
    db.flush()
    for a in rows:
        record_tables(db, db.get(Reservation, a.reservation_id), user, before.get(a.reservation_id, []))
    db.commit()
    return {"ok": True}


# ---------- CSV取込(宿泊者リスト) ----------
IMPORT_MAX_BYTES = 5 * 1024 * 1024
MANUAL_ACTIONS = ("create", "update", "delete", "restore")  # CSV取込以外の操作
KIND_LABELS = {"reservations": "①予約", "rooms": "②部屋割り"}


class ImportFileIn(BaseModel):
    name: str = Field(max_length=255)
    data: str  # base64


class ImportIn(BaseModel):
    files: list[ImportFileIn] = Field(max_length=10)


def import_user(user: User = Depends(current_user)) -> User:
    if user.role not in IMPORT_ROLES:
        raise HTTPException(403, "CSVの取り込み権限がありません")
    return user


def read_import_files(body: ImportIn) -> tuple[list[dict], dict[str, list[list[str]]]]:
    """ファイルごとに種別を判定する。同じ種別が複数あれば後のファイルを使う"""
    files, rows_by_kind = [], {}
    for f in body.files:
        try:
            raw = base64.b64decode(f.data, validate=True)
        except (binascii.Error, ValueError):
            raise HTTPException(400, f"{f.name} を読み込めません")
        if len(raw) > IMPORT_MAX_BYTES:
            raise HTTPException(400, f"{f.name} が大きすぎます(5MBまで)")
        try:
            rows = importer.read_rows(raw)
        except importer.CsvError as e:
            raise HTTPException(400, f"{f.name}: {e}")
        kind = importer.detect_kind(rows)
        files.append({"name": f.name, "kind": kind, "kind_label": KIND_LABELS.get(kind, "判別できません"),
                      "rows": len(rows)})
        if kind:
            rows_by_kind[kind] = rows
    return files, rows_by_kind


def import_plan(db: Session, stays: list[importer.Stay]) -> dict:
    """取り込み内容を計算する(DBは変更しない)"""
    keys = [s.key for s in stays]
    existing = {}
    for r in db.scalars(select(Reservation).where(Reservation.ext_key.in_(keys))) if keys else []:
        existing[(r.meal, r.ext_key, r.date)] = r
    creates, updates, alerts = [], [], []
    unchanged = skipped_deleted = 0
    seen = set()
    today = now_jst().date()  # 過去の日付は登録も要確認もしない(長期滞在の過去分など)
    for s in stays:
        targets = {"dinner": s.dinner, "breakfast": s.breakfast}
        for meal in MEALS:
            wanted = [] if s.cancelled or not targets[meal] else s.dates(meal)
            for i, d in enumerate(wanted):
                if d < today:
                    continue  # 何泊目・泊数は滞在全体で数える
                seen.add((meal, s.key, d))
                r = existing.get((meal, s.key, d))
                values = {"room": s.room, "guest_name": s.name}
                if r is None:
                    creates.append((s, meal, d, i + 1, len(wanted)))
                elif r.deleted_at is not None:
                    skipped_deleted += 1  # 手で削除したものは作り直さない
                elif any(getattr(r, k) != v for k, v in values.items()):
                    updates.append((r, values))
                else:
                    unchanged += 1
            # 取込済みだが、取消・対象外・日程外になったもの(自動では消さずに知らせる)
            stale = [r for (m, k, d), r in existing.items()
                     if m == meal and k == s.key and r.deleted_at is None and d >= today and (m, k, d) not in seen]
            if stale:
                reason = ("取消になりました" if s.cancelled
                          else f"{MEAL_LABELS[meal]}の対象外になりました" if not targets[meal]
                          else "日程から外れました")
                stale.sort(key=lambda r: r.date)
                manual = set(db.scalars(select(ReservationHistory.reservation_id).where(
                    ReservationHistory.reservation_id.in_([r.id for r in stale]),
                    ReservationHistory.action.in_(MANUAL_ACTIONS))))
                alerts.append({
                    "key": s.key, "meal": meal, "meal_label": MEAL_LABELS[meal], "reason": reason,
                    "room": stale[0].room, "guest_name": stale[0].guest_name,
                    "dates": [r.date.isoformat() for r in stale], "manual": bool(manual),
                    "link": f"/{meal}?d={stale[0].date.isoformat()}&hl={stale[0].id}",
                })
    by_slot = {}
    for s, meal, *_ in creates:
        if meal == "dinner":
            by_slot[s.dinner_time or ""] = by_slot.get(s.dinner_time or "", 0) + 1
    return {
        "creates": creates, "updates": updates, "alerts": alerts,
        "summary": {
            "stays": len(stays),
            "cancelled": sum(s.cancelled for s in stays),
            "not_target": sum(not s.cancelled and not (s.dinner or s.breakfast) for s in stays),
            "unassigned": sum(not s.cancelled and (s.dinner or s.breakfast) and not s.rooms for s in stays),
            "dinner": sum(m == "dinner" for _, m, *_ in creates),
            "dinner_by_slot": [{"slot": k, "count": v}
                               for k, v in sorted(by_slot.items(), key=lambda x: x[0] or "99")],
            "breakfast": sum(m == "breakfast" for _, m, *_ in creates),
            "update": len(updates), "unchanged": unchanged, "skipped_deleted": skipped_deleted,
        },
    }


def prepare_import(body: ImportIn, db: Session) -> tuple[list[dict], dict | None]:
    files, rows = read_import_files(body)
    if not ("reservations" in rows and "rooms" in rows):
        return files, None
    try:
        stays = importer.build_stays(rows["reservations"], rows["rooms"])
    except importer.CsvError as e:
        raise HTTPException(400, str(e))
    return files, import_plan(db, stays)


def import_response(files: list[dict], plan: dict | None) -> dict:
    kinds = {f["kind"] for f in files}
    return {"files": files, "ready": plan is not None,
            "missing": [label for kind, label in KIND_LABELS.items() if kind not in kinds],
            "summary": plan["summary"] if plan else None, "alerts": plan["alerts"] if plan else []}


@app.post("/api/import/preview")
def import_preview(body: ImportIn, _: User = Depends(import_user), db: Session = Depends(get_db)):
    return import_response(*prepare_import(body, db))


@app.post("/api/import/commit")
def import_commit(body: ImportIn, user: User = Depends(import_user), db: Session = Depends(get_db)):
    files, plan = prepare_import(body, db)
    if plan is None:
        raise HTTPException(400, "①予約と②部屋割りの2つのファイルがそろっていないため、取り込みを中止しました")
    now = now_jst()
    stay_ids = {}
    for s, meal, d, night_no, nights in plan["creates"]:
        stay_id = stay_ids.setdefault((s.key, meal), uuid.uuid4().hex) if nights > 1 else None
        r = Reservation(meal=meal, date=d, nights=nights, night_no=night_no, stay_id=stay_id, ext_key=s.key,
                        room=s.room, guest_name=s.name, adults=0, children=0, infants=0,
                        time_slot=s.dinner_time if meal == "dinner" else None, allergy="", note="",
                        created_at=now, created_by=user.id, updated_at=now, updated_by=user.id)
        db.add(r)
        db.flush()
        record(db, r, user, "import", {f: [None, v] for f, v in snapshot(r).items()})
    for r, values in plan["updates"]:
        change(db, r, values, user, action="import_update")
    db.commit()
    return import_response(files, plan)


# ---------- ユーザー管理 ----------
def user_dict(u: User) -> dict:
    return {"id": u.id, "username": u.username, "display_name": u.display_name,
            "role": u.role, "active": u.active}


Role = Literal["developer", "admin", "front", "restaurant"]


def set_role(u: User, role: str) -> None:
    u.role, u.is_admin = role, role in ADMIN_ROLES


def check_role_change(me_: User, target_role: str | None, target: User | None = None) -> None:
    """最上位ロールのアカウントを作る・編集する・最上位ロールにするのは最上位ロールのみ"""
    if me_.role == "developer":
        return
    if target_role == "developer" or (target and target.role == "developer"):
        raise HTTPException(403, f"{ROLES['developer']} のアカウントは {ROLES['developer']} のみ操作できます")


class UserCreateIn(BaseModel):
    username: str = Field(pattern=r"^[A-Za-z0-9_.-]{2,64}$")
    display_name: str = Field(min_length=1, max_length=64)
    password: str = Field(min_length=4, max_length=128)
    role: Role = "front"


class UserUpdateIn(BaseModel):
    display_name: str | None = Field(default=None, min_length=1, max_length=64)
    password: str | None = Field(default=None, min_length=4, max_length=128)
    role: Role | None = None
    active: bool | None = None


@app.get("/api/users")
def list_users(_: User = Depends(admin_user), db: Session = Depends(get_db)):
    return [user_dict(u) for u in db.scalars(select(User).order_by(User.id))]


@app.post("/api/users")
def create_user(body: UserCreateIn, me_: User = Depends(admin_user), db: Session = Depends(get_db)):
    check_role_change(me_, body.role)
    if db.scalar(select(User).where(User.username == body.username)):
        raise HTTPException(400, "そのログインIDは既に使われています")
    u = User(username=body.username, display_name=body.display_name.strip(),
             password_hash=hash_password(body.password))
    set_role(u, body.role)
    db.add(u)
    db.commit()
    return user_dict(u)


@app.put("/api/users/{uid}")
def update_user(uid: int, body: UserUpdateIn, me_: User = Depends(admin_user), db: Session = Depends(get_db)):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "見つかりません")
    check_role_change(me_, body.role, u)
    if u.id == me_.id and ((body.role and body.role != me_.role) or body.active is False):
        raise HTTPException(400, "自分自身のロール変更・無効化はできません")
    if body.display_name is not None:
        u.display_name = body.display_name.strip()
    if body.password is not None:
        u.password_hash = hash_password(body.password)
    if body.role is not None:
        set_role(u, body.role)
    if body.active is not None:
        u.active = body.active
    db.commit()
    return user_dict(u)


# ---------- チャット ----------
CHAT_LIMIT = 200


class ChatIn(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)

    body: str = Field(default="", max_length=2000)
    reservation_id: int | None = None


class ReadIn(BaseModel):
    last_id: int


def chat_dict(m: ChatMessage, users: dict[int, User], res: dict[int, Reservation]) -> dict:
    u = users.get(m.user_id)
    r = res.get(m.reservation_id)
    retracted = m.retracted_at is not None
    return {
        "id": m.id, "user_id": m.user_id, "name": u.display_name if u else "",
        "role": m.role, "role_label": ROLES.get(m.role, ""),
        "body": "" if retracted else m.body, "retracted": retracted,
        "created_at": iso(m.created_at),
        "reservation": None if retracted or not r else {
            "id": r.id, "meal": r.meal, "date": r.date.isoformat(), "room": r.room,
            "guest_name": r.guest_name, "time_slot": r.time_slot, "deleted": r.deleted_at is not None},
    }


def unread_count(db: Session, user: User) -> int:
    last = db.scalar(select(ChatRead.last_read_id).where(ChatRead.user_id == user.id)) or 0
    return db.scalar(select(func.count()).select_from(ChatMessage).where(
        ChatMessage.id > last, ChatMessage.user_id != user.id, ChatMessage.retracted_at.is_(None))) or 0


@app.get("/api/chat/messages")
def chat_messages(user: User = Depends(current_user), db: Session = Depends(get_db)):
    msgs = list(db.scalars(select(ChatMessage).order_by(ChatMessage.id.desc()).limit(CHAT_LIMIT)))[::-1]
    users = {u.id: u for u in db.scalars(select(User))}
    rids = {m.reservation_id for m in msgs if m.reservation_id}
    res = {r.id: r for r in db.scalars(select(Reservation).where(Reservation.id.in_(rids)))} if rids else {}
    reads = [{"user_id": uid, "name": users[uid].display_name, "role_label": ROLES.get(users[uid].role, ""),
              "last_read_id": last}
             for uid, last in db.execute(select(ChatRead.user_id, ChatRead.last_read_id))
             if uid in users and users[uid].active]
    return {"me": user.id, "messages": [chat_dict(m, users, res) for m in msgs], "reads": reads}


@app.post("/api/chat/messages")
def post_chat(body: ChatIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not body.body and not body.reservation_id:
        raise HTTPException(400, "メッセージを入力してください")
    r = db.get(Reservation, body.reservation_id) if body.reservation_id else None
    if body.reservation_id and not r:
        raise HTTPException(400, "添付した予約が見つかりません")
    m = ChatMessage(user_id=user.id, role=user.role, body=body.body, reservation_id=body.reservation_id)
    db.add(m)
    db.flush()
    mark_read(db, user, m.id)
    db.commit()
    return {"id": m.id}


@app.post("/api/chat/messages/{mid}/retract")
def retract_chat(mid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    m = db.get(ChatMessage, mid)
    if not m or m.user_id != user.id:
        raise HTTPException(404, "取り消せるメッセージが見つかりません")
    if m.retracted_at is None:
        m.retracted_at = now_jst()
        db.commit()
    return {"ok": True}


def mark_read(db: Session, user: User, last_id: int) -> None:
    cr = db.get(ChatRead, user.id)
    if not cr:
        db.add(ChatRead(user_id=user.id, last_read_id=last_id))
    elif last_id > cr.last_read_id:
        cr.last_read_id = last_id


@app.post("/api/chat/read")
def chat_read(body: ReadIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    latest = db.scalar(select(func.max(ChatMessage.id))) or 0
    mark_read(db, user, min(body.last_id, latest))
    db.commit()
    return {"unread": unread_count(db, user)}


@app.get("/api/chat/unread")
def chat_unread(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return {"unread": unread_count(db, user)}


# ---------- 集計 ----------
MEAL_LABELS = {"dinner": "夕食", "breakfast": "朝食"}
WEEKDAYS = "月火水木金土日"
# (キー, 見出し) — 画面と Excel の列順
SUMMARY_COLS = [
    ("groups", "組数"), ("adults", "大人"), ("children", "幼児"), ("infants", "席のみ"), ("total", "計"),
    ("entered", "入場済(組)"),
]
SUMMED = ("adults", "children", "infants")


def summarize(db: Session, meal: str, start: date, end: date) -> dict:
    if end < start:
        raise HTTPException(400, "終了日は開始日以降にしてください")
    if (end - start).days > 366:
        raise HTTPException(400, "期間は1年以内にしてください")
    days = {start + timedelta(days=i): dict.fromkeys((k for k, _ in SUMMARY_COLS), 0)
            for i in range((end - start).days + 1)}
    rows = db.scalars(select(Reservation).where(
        Reservation.meal == meal, Reservation.date >= start, Reservation.date <= end,
        Reservation.deleted_at.is_(None)))  # 削除済みは集計しない
    for r in rows:
        a = days[r.date]
        a["groups"] += 1
        for k in SUMMED:
            a[k] += getattr(r, k)
        a["total"] += r.adults + r.children  # 席のみは計に含めない
        a["entered"] += r.entered_at is not None
    total = {k: sum(a[k] for a in days.values()) for k, _ in SUMMARY_COLS}
    return {
        "meal": meal, "start": start.isoformat(), "end": end.isoformat(),
        "columns": [{"key": k, "label": l} for k, l in SUMMARY_COLS],
        "days": [{"date": d.isoformat(), "weekday": WEEKDAYS[d.weekday()], **a} for d, a in days.items()],
        "total": total,
    }


@app.get("/api/{meal}/summary")
def meal_summary(meal: str, start: date, end: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    return summarize(db, check_meal(meal), start, end)


@app.get("/api/{meal}/summary.xlsx")
def meal_summary_xlsx(meal: str, start: date, end: date, _: User = Depends(current_user),
                      db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    data = summarize(db, check_meal(meal), start, end)
    title = f"{MEAL_LABELS[meal]}集計"
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append([f"{title}　{start:%Y/%m/%d}〜{end:%Y/%m/%d}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"出力日時 {now_jst():%Y/%m/%d %H:%M}(削除済みの予約は除く)"])
    ws["A2"].font = Font(size=9, color="777777")
    ws.append([])
    head = ["日付", "曜日"] + [l for _, l in SUMMARY_COLS]
    ws.append(head)
    thin = Side(style="thin", color="BBBBBB")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E6F1FB")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    # 合計行は見出しの直後(5行目)
    ws.append(["合計", ""] + [data["total"][k] for k, _ in SUMMARY_COLS])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F2F2F2")
        c.border = Border(top=thin, left=thin, right=thin, bottom=Side(style="medium", color="8FA9C4"))
    for d in data["days"]:
        ws.append([date.fromisoformat(d["date"]), d["weekday"]] + [d[k] for k, _ in SUMMARY_COLS])
        row = ws[ws.max_row]
        row[0].number_format = "yyyy/mm/dd"
        color = {"土": "185FA5", "日": "C0392B"}.get(d["weekday"])
        for c in row:
            c.border = border
            if color and c.column <= 2:
                c.font = Font(color=color)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 5
    for i in range(3, len(head) + 1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = 11
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "C6"  # 見出しと合計行を固定
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:5"  # 印刷時は各ページに見出しと合計行

    buf = BytesIO()
    wb.save(buf)
    fname = f"{title}_{start:%Y%m%d}-{end:%Y%m%d}.xlsx"
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ---------- 操作ログ ----------
FIELD_LABELS = {
    "date": "日付", "nights": "泊数", "night_no": "何泊目", "time_slot": "時間", "room": "部屋",
    "guest_name": "代表者名", "adults": "大人", "children": "幼児", "infants": "席のみ",
    "allergy": "アレルギー", "note": "備考", "group_id": "グループ", "entered_at": "ステータス",
    "tables": "テーブル",
}
RES_ACTIONS = {"create": "登録", "update": "変更", "delete": "削除", "restore": "復元",
               "import": "CSV取込", "import_update": "CSV取込(更新)"}
AUTH_ACTIONS = {"login": "ログイン", "login_failed": "ログイン失敗", "logout": "ログアウト",
                "password_change": "パスワード変更"}
LOG_MAX_DAYS = 93


def fmt_value(field: str, v) -> str:
    if field == "time_slot":
        return v or "未定"
    if field == "group_id":
        return "あり" if v else "なし"
    if field == "entered_at":
        return "入場済" if v else "空白"
    if field == "tables":
        return "・".join(v) if v else "なし"
    return "(空欄)" if v in (None, "") else str(v)


def change_detail(action: str, changes: dict) -> str:
    if action in ("create", "import"):  # 登録時は入力された項目だけ
        return "、".join(f"{FIELD_LABELS.get(f, f)} {fmt_value(f, b)}" for f, (_, b) in changes.items()
                        if b not in (None, "", 0) and f not in ("date", "nights", "night_no", "room", "guest_name"))
    return "、".join(f"{FIELD_LABELS.get(f, f)}: {fmt_value(f, a)} → {fmt_value(f, b)}" for f, (a, b) in changes.items())


@app.get("/api/logs")
def operation_logs(start: date, end: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    """予約・チャット・ログインの操作を新しい順にまとめて返す"""
    if end < start:
        raise HTTPException(400, "終了日は開始日以降にしてください")
    if (end - start).days >= LOG_MAX_DAYS:
        raise HTTPException(400, f"期間は{LOG_MAX_DAYS}日以内にしてください")
    t0 = datetime.combine(start, datetime.min.time())
    t1 = datetime.combine(end + timedelta(days=1), datetime.min.time())
    users = {u.id: u for u in db.scalars(select(User))}

    def who(uid, fallback=""):
        u = users.get(uid)
        return {"user": u.display_name if u else fallback, "role": u.role if u else "",
                "role_label": ROLES.get(u.role, "") if u else ""}

    logs = []
    # 予約
    for h, r in db.execute(select(ReservationHistory, Reservation)
                           .join(Reservation, Reservation.id == ReservationHistory.reservation_id)
                           .where(ReservationHistory.changed_at >= t0, ReservationHistory.changed_at < t1)):
        logs.append({
            "_id": h.id, "at": iso(h.changed_at), "kind": "reservation", "kind_label": "予約",
            "action": "入場済" if set(h.changes) == {"entered_at"} else RES_ACTIONS.get(h.action, h.action),
            "target": f"{MEAL_LABELS.get(r.meal, '')} {r.date:%m/%d} {r.room} {r.guest_name}".strip(),
            "link": f"/{r.meal}?d={r.date.isoformat()}&hl={r.id}",
            "detail": change_detail(h.action, h.changes or {}), **who(h.changed_by)})
    # チャット(取り消されたメッセージの本文はログにも出さない)
    chat_q = select(ChatMessage).where(
        ((ChatMessage.created_at >= t0) & (ChatMessage.created_at < t1)) |
        ((ChatMessage.retracted_at >= t0) & (ChatMessage.retracted_at < t1)))
    for m in db.scalars(chat_q):
        base = {"_id": m.id, "kind": "chat", "kind_label": "チャット", "target": "チャット", "link": "/chat", **who(m.user_id)}
        body = "(取り消し済み)" if m.retracted_at else m.body
        if t0 <= m.created_at < t1:
            logs.append({**base, "at": iso(m.created_at), "action": "投稿", "detail": body})
        if m.retracted_at and t0 <= m.retracted_at < t1:
            logs.append({**base, "_id": m.id + 0.5, "at": iso(m.retracted_at), "action": "取り消し", "detail": ""})
    # ログイン
    for a in db.scalars(select(AuthLog).where(AuthLog.at >= t0, AuthLog.at < t1)):
        logs.append({"_id": a.id, "at": iso(a.at), "kind": "auth", "kind_label": "ログイン",
                     "action": AUTH_ACTIONS.get(a.action, a.action), "target": "", "link": "",
                     "detail": f"ログインID: {a.username}" if a.action == "login_failed" else "",
                     "failed": a.action == "login_failed", **who(a.user_id, a.username)})
    # 同じ秒の操作は記録順(各テーブルのID順)で並べる
    logs.sort(key=lambda x: (x["at"], x.pop("_id")), reverse=True)
    return logs


# ---------- ページ ----------
PAGES = {"dinner", "tables", "breakfast", "dinner-summary", "breakfast-summary", "chat", "logs", "admin"}


@app.get("/favicon.ico", include_in_schema=False)
def favicon():
    return FileResponse(BASE / "static" / "favicon.ico", media_type="image/x-icon")


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.get("/")
def root():
    return RedirectResponse("/dinner")


@app.get("/login")
def login_page(request: Request, db: Session = Depends(get_db)):
    if session_user(request, db):
        return RedirectResponse("/dinner")
    return FileResponse(BASE / "pages" / "login.html")


@app.get("/{page}")
def page(page: str, request: Request, db: Session = Depends(get_db)):
    if page not in PAGES:
        return RedirectResponse("/dinner")
    user = session_user(request, db)
    if not user:
        return RedirectResponse("/login")
    if page == "admin" and user.role not in ADMIN_ROLES:
        return RedirectResponse("/dinner")
    return FileResponse(BASE / "pages" / f"{page.replace('-', '_')}.html")
